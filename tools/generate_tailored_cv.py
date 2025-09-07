import json
from openai import OpenAI
from docxtpl import DocxTemplate, RichText
from docx import Document
import shutil
import tempfile
import os

def load_static_cv_data(cv_data: dict, projects_info_file_path: str):
    """
    Load static data associated with dynamically selected CV components
    For example, the LLM chooses which projects to include, but the dates associated with these projects are static
    The static components are inserted independently to reduce errors
    """
    with open(projects_info_file_path, "r", encoding="utf-8") as f:
        projects = json.load(f)

    for i, project in enumerate(cv_data["Relevant Projects"]):
        project_title = project["Title"]
        if project_title in projects:
            details = projects[project_title]
            cv_data["Relevant Projects"][i]["Dates"] = details.get("Dates", "")
            cv_data["Relevant Projects"][i]["Link"] = details.get("Link", "")
            cv_data["Relevant Projects"][i]["Description"] = details.get("Description", "")


def preprocess_template(template_file_path: str) -> str:
    """
    Replace literal '&' characters in the template with '&amp;' to survive docxtpl rendering.
    Saves to a temporary file and returns the path.
    """
    temp_dir = tempfile.mkdtemp()
    temp_template_path = os.path.join(temp_dir, "preprocessed_template.docx")
    shutil.copy(template_file_path, temp_template_path)

    doc = Document(temp_template_path)
    for p in doc.paragraphs:
        for run in p.runs:
            run.text = run.text.replace("&", "&amp;")
    doc.save(temp_template_path)
    return temp_template_path


def render_cv_to_docx(cv_data: dict, projects_info_file_path: str, template_file_path: str, output_file_path: str):
    """
    Insert LLM-generated tailored CV components into template CV
    Returns:
        str: Path to the output CV
    """
    load_static_cv_data(cv_data, projects_info_file_path)

    # Preprocess template to preserve ampersands in static text
    preprocessed_template_path = preprocess_template(template_file_path)
    doc = DocxTemplate(preprocessed_template_path)

    # Build RichText hyperlinks for the two projects
    project_1_link = RichText()
    project_1_link.add("View Project", url_id=doc.build_url_id(cv_data["Relevant Projects"][0]["Link"]))

    project_2_link = RichText()
    project_2_link.add("View Project", url_id=doc.build_url_id(cv_data["Relevant Projects"][1]["Link"]))

    # Rename keys for Jinja (no spaces, consistent with template placeholders)
    context = {
        "Profile": cv_data["Profile"],
        "Key_Competencies": cv_data["Technical Skills"]["Key Competencies"],
        "Programming_Languages": cv_data["Technical Skills"]["Programming Languages"],
        "Frameworks_and_Libraries": cv_data["Technical Skills"]["Frameworks & Libraries"],
        "Tools_and_Platforms": cv_data["Technical Skills"]["Tools & Platforms"],
        "Project_1_Title": cv_data["Relevant Projects"][0]["Title"],
        "Project_1_Dates": cv_data["Relevant Projects"][0]["Dates"],
        "Project_1_Skills": cv_data["Relevant Projects"][0]["Skills"],
        "Project_1_Link": project_1_link,
        "Project_1_Description": cv_data["Relevant Projects"][0]["Description"],
        "Project_2_Title": cv_data["Relevant Projects"][1]["Title"],
        "Project_2_Dates": cv_data["Relevant Projects"][1]["Dates"],
        "Project_2_Skills": cv_data["Relevant Projects"][1]["Skills"],
        "Project_2_Link": project_2_link,
        "Project_2_Description": cv_data["Relevant Projects"][1]["Description"],
    }

    print("Starting attempt")
    doc.render(context)
    doc.save(output_file_path)
    print("Attempt complete")
    print(f"CV saved to {output_file_path}")
    return output_file_path



def generate_tailored_cv(path_to_job_listing: str, path_to_projects_info: str, path_to_response_format: str, path_to_cv_template: str, path_to_output_cv: str) -> str:
    """
    Generate a tailored CV for job listing.
    The CV is stored in a word document which the user can inspect and edit.
    If they deem it to be acceptable, they can request for it to be converted to a pdf.
    """
    try:
        client = OpenAI()

        with open(path_to_job_listing, "r", encoding="utf-8") as f:
            job_listing = f.read()

        with open(path_to_response_format, "r", encoding="utf-8") as f:
            response_format = json.load(f)

        response = client.responses.create(
            input=f"""
            =====================================================================
            Job Listing:

            {job_listing}
            =====================================================================
            
            Return a tailored CV for this job listing as a JSON object with fields:
            Profile, Technical Skills, Relevant Projects.
            """,
            prompt={
                "id": "pmpt_68b0c358291c81968a00e9414e386276009a458f13092457",
                "version": "13"
            },
            text=response_format  # schema for structured model output
        )

        cv_data = json.loads(response.output_text)
        print(json.dumps(cv_data, indent=2))
        output_file_path = render_cv_to_docx(cv_data, path_to_projects_info, path_to_cv_template, path_to_output_cv)

        return output_file_path

    except Exception as e:
        print(f"Error generating CV: {e}")
        return ""
